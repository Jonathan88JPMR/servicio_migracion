# -*- coding: utf-8 -*-
"""
Genera la documentación del proyecto en Word (.docx) y el Excel de endpoints/SPs.

Salida: docs/word/
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(BASE, "docs")
IMG = os.path.join(DOCS, "img")
OUT = os.path.join(DOCS, "word")
FECHA = "24/09/2026"
AUTOR = "Jonathan Marquina"

os.makedirs(OUT, exist_ok=True)

# ---------------------------------------------------------------------------
# Conversor Markdown -> docx
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def add_runs(par, text):
    """Agrega texto con **negrita** y `codigo` inline."""
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        else:
            par.add_run(tok)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def code_paragraph(doc, line):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run(line if line else " ")
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    return p


def is_table_line(line):
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator_row(cells):
    return all(re.fullmatch(r":?-{2,}:?", c or "---") for c in cells)


def md_to_docx(md_text, doc, img_dir):
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # bloque de código
        if stripped.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_paragraph(doc, lines[i].rstrip("\n"))
                i += 1
            i += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # tabla
        if is_table_line(line) and i + 1 < len(lines) and is_table_line(lines[i + 1]) \
                and is_separator_row(split_row(lines[i + 1])):
            header = split_row(line)
            i += 2
            rows = []
            while i < len(lines) and is_table_line(lines[i]):
                rows.append(split_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                r = p.add_run(txt.replace("**", ""))
                r.bold = True
                r.font.size = Pt(9)
                shade_cell(cell, "D9E2F3")
            for ri, row in enumerate(rows):
                for c in range(len(header)):
                    cell = table.rows[ri + 1].cells[c]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    txt = row[c] if c < len(row) else ""
                    add_runs(p, txt)
                    for r in p.runs:
                        if r.font.size is None:
                            r.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # imagen
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if m:
            alt, rel = m.group(1), m.group(2)
            path = os.path.normpath(os.path.join(img_dir, rel))
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(path, width=Inches(6.3))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(f"Figura: {alt}")
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            else:
                doc.add_paragraph(f"[Imagen no encontrada: {rel}]")
            i += 1
            continue

        # encabezados
        hm = re.match(r"^(#{1,5})\s+(.*)$", stripped)
        if hm:
            level = min(len(hm.group(1)), 4)
            doc.add_heading(hm.group(2).replace("**", ""), level=level)
            i += 1
            continue

        # separador horizontal
        if re.fullmatch(r"-{3,}|_{3,}", stripped):
            i += 1
            continue

        # listas
        lm = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if lm:
            indent = len(lm.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
            add_runs(p, lm.group(2))
            i += 1
            continue
        nm = re.match(r"^(\s*)\d+[.)]\s+(.*)$", line)
        if nm:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, nm.group(2))
            i += 1
            continue

        # párrafo normal
        if stripped:
            p = doc.add_paragraph()
            add_runs(p, stripped)
        i += 1


def add_portada(doc, titulo, subtitulo, extra=None):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitulo)
    r2.font.size = Pt(15)
    r2.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    if extra:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(extra)
        r3.font.size = Pt(11)
        r3.italic = True
    doc.add_paragraph()
    for label, val in (("Proyecto", "Migración de documentos a Spring Gestión Documental"),
                       ("Elaborado por", AUTOR), ("Fecha", FECHA)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(val)
    doc.add_page_break()


def convert_md(md_path, out_name, titulo, subtitulo):
    doc = Document()
    add_portada(doc, titulo, subtitulo)
    with open(md_path, encoding="utf-8") as f:
        md_to_docx(f.read(), doc, os.path.dirname(md_path))
    out = os.path.join(OUT, out_name)
    doc.save(out)
    print("OK", out_name)


# ---------------------------------------------------------------------------
# Extracción de secciones (## N.) de un md
# ---------------------------------------------------------------------------

def extract_sections(md_text, numbers):
    """Devuelve el texto de las secciones '## N.' cuyo número está en `numbers`."""
    parts = re.split(r"(?m)^(## )", md_text)
    out = []
    i = 1
    while i < len(parts):
        if parts[i] == "## ":
            header = parts[i + 1].splitlines()[0]
            m = re.match(r"(\d+)\.", header.strip())
            if m and int(m.group(1)) in numbers:
                out.append(parts[i] + parts[i + 1])
            i += 2
        else:
            i += 1
    return "\n".join(out)


def build_component_doc(out_name, titulo, subtitulo, intro_md, body_md):
    doc = Document()
    add_portada(doc, titulo, subtitulo)
    md_to_docx(intro_md, doc, DOCS)
    md_to_docx(body_md, doc, DOCS)
    out = os.path.join(OUT, out_name)
    doc.save(out)
    print("OK", out_name)


def main():
    doc04 = open(os.path.join(DOCS, "04_DOCUMENTACION_CODIGO.md"), encoding="utf-8").read()

    # ----- 1. Documentación por proyecto -----
    intro_srv = """# Documentación técnica — `servicio_documentos`

Servicio migrador de adjuntos (Python 3, empaquetado con PyInstaller) más el
dashboard legado (Flask + waitress + Bootstrap) que lo supervisa. Migra los
documentos del portal de proveedores hacia Spring Gestión Documental conservando
la relación con obligación, orden de compra/servicio, proveedor, empresa y grupo
documental.

| Dato | Detalle |
|---|---|
| Carpeta | `servicio_documentos/` |
| Lenguaje | Python 3 (`pyodbc`, `Flask`, `waitress`) |
| Ejecutables | `migrar_prueba.exe` / `migrar_produccion.exe` (`build_servicio.bat`) |
| Programa principal | `migrar_adjuntos_spring_prueba.py` (prueba) / `migrar_adjuntos_spring.py` (producción) |
| Dashboard legado | `frontend/app.py` + `frontend/server.py` (waitress) |
| Programación | Tarea de Windows cada 2 horas (`crear_tarea.ps1`) |
"""
    build_component_doc("DOC_Servicio_Documentos.docx",
                        "Servicio de Migración de Documentos",
                        "servicio_documentos — servicio migrador Python + dashboard Flask",
                        intro_srv, extract_sections(doc04, {1, 2, 3, 4, 5}))

    intro_api = """# Documentación técnica — `api_migracion_documentos`

API REST en .NET 8 con arquitectura limpia (Domain / Application /
Infraestructure). **No migra archivos**: orquesta el dashboard — dispara el
servicio migrador, lee su log y configuración, y consulta las bases vía stored
procedures `sp_mig_*` con parámetro `@json` que devuelven `FOR JSON`.

| Dato | Detalle |
|---|---|
| Carpeta | `api_migracion_documentos/` |
| Framework | .NET 8 Web API (`Microsoft.Data.SqlClient`, `ClosedXML`) |
| Controlador | `MigracionController` — ruta base `api/migracion-documentos` |
| Convención | POST + body JSON; GET solo para previews binarios |
| Despliegue | IIS (`api_migracion_doc`, publicación `dotnet publish -c Release`) |
| URL producción | `https://apimigraciondoc.agroapps.net:7631` |
"""
    build_component_doc("DOC_API_Migracion_Documentos.docx",
                        "API de Migración de Documentos",
                        "api_migracion_documentos — Web API .NET 8 (clean architecture)",
                        intro_api, extract_sections(doc04, {6, 8}))

    intro_front = """# Documentación técnica — `migracion_documentos`

Frontend web del sistema de migración. SPA Angular 20 con componentes
standalone, rutas con hash (`withHashLocation`, evita URL Rewrite en IIS),
Bootstrap 5 + PrimeNG (`p-chart`) + sweetalert2.

| Dato | Detalle |
|---|---|
| Carpeta | `migracion_documentos/` |
| Framework | Angular 20 + TypeScript 5.9 |
| Librerías | Bootstrap 5.3, PrimeNG 20, Chart.js 4, sweetalert2, RxJS 7 |
| Servicios | `MigracionService` (llamadas POST), `AlertService`, `RunStateService` |
| API local | `http://localhost:5080/api/migracion-documentos` (`environment.ts`) |
| API producción | `https://apimigraciondoc.agroapps.net:7631/api/migracion-documentos` |
| Build | `ng build` → `dist/` → `wwwroot` de la API en IIS |
"""
    extra_front = """
## Servicios

| Servicio | Archivo | Función |
|---|---|---|
| `MigracionService` | `modules/main/services/migracion.service.ts` | Encapsula todos los POST a la API con `lastValueFrom`: `entornos`, `cambiarEntorno`, `configGet`/`configUpdate`, `stats`, `archivos`, `archivosExcel` (blob), `urlPreviewAp`/`urlPreviewMs`, `log`, `run`, `runAll`, `runOrden`, `runRango`, `runStatus` |
| `RunStateService` | `modules/main/services/run-state.service.ts` | Estado compartido de la ejecución activa entre "Ejecución" y "Ejecución personalizada": `activeRunId`, poll cada 2 s a `run-status`, temporizador y consola de salida |
| `AlertService` | `shared/alertas/alerts.service.ts` | Centraliza los diálogos de sweetalert2 (confirmaciones, errores, avisos) |

## Configuración de la aplicación

- `app.config.ts`: `provideRouter(routes, withHashLocation())`, `provideHttpClient()`, animaciones y PrimeNG con tema **Aura**.
- `app.routes.ts`: `LayoutComponent` como shell con menú lateral; cada página se carga con `loadComponent` (lazy).
- `environments/environment.ts` → `apiUrl` local (`http://localhost:5080/api/migracion-documentos`).
- `environments/environment.prod.ts` → `apiUrl` productivo (`https://apimigraciondoc.agroapps.net:7631/api/migracion-documentos`).

## Estructura de carpetas

```
migracion_documentos/
├── src/
│   ├── index.html
│   ├── main.ts                         bootstrap con appConfig
│   ├── styles.css                      Bootstrap + estilos globales
│   ├── environments/                   apiUrl por ambiente
│   └── app/
│       ├── app.config.ts               providers (router hash, HttpClient, PrimeNG)
│       ├── app.routes.ts               rutas lazy con LayoutComponent
│       ├── app.ts                      componente raíz
│       ├── shared/alertas/             AlertService (sweetalert2)
│       └── modules/main/
│           ├── services/               MigracionService, RunStateService
│           └── pages/                  dashboard, archivos, configuracion,
│                                     ejecucion, ejecucion-personalizada,
│                                     ejecucion-rango, log, layout
├── angular.json / package.json / tsconfig*.json
└── dist/                               build de producción (ng build)
```

## Build y despliegue

- Desarrollo: `ng serve` (proxy a la API local en `:5080`).
- Producción: `ng build` → `dist/`; el contenido se copia al `wwwroot` del paquete
  `api_migracion_doc` en IIS. La API sirve el SPA con `UseDefaultFiles` +
  `UseStaticFiles` + `MapFallbackToFile("index.html")`; las rutas usan hash
  (`#/dashboard`, `#/archivos`, …) por lo que no se requiere URL Rewrite.
"""
    build_component_doc("DOC_Frontend_Migracion_Documentos.docx",
                        "Frontend Web de Migración de Documentos",
                        "migracion_documentos — SPA Angular 20 + PrimeNG + Bootstrap",
                        intro_front, extract_sections(doc04, {7}) + extra_front)

    # ----- 2. Conversión de todos los documentos md -----
    conversiones = [
        ("00_INDICE.md", "00_Indice_Documentacion.docx", "Índice de documentación", "Proyecto Migración de Documentos"),
        ("01_MANUAL_DESPLIEGUE.md", "01_Manual_Despliegue.docx", "Manual de despliegue en el servidor", "Servicio, tarea programada, dashboard e IIS"),
        ("02_MANUAL_USUARIO.md", "02_Manual_Usuario.docx", "Manual de usuario", "Dashboard de Migración de Adjuntos"),
        ("03_PROCEDIMIENTO.md", "03_Procedimiento_Operativo.docx", "Procedimiento operativo", "Proceso del aplicativo: monitoreo, errores y paso a producción"),
        ("04_DOCUMENTACION_CODIGO.md", "04_Documentacion_Codigo.docx", "Documentación del código", "Servicio, APIs, frontends, scripts y tablas"),
        ("05_FLUJO_COMPLETO.md", "05_Flujo_Completo.docx", "Flujo completo del servicio", "Origen → resolución → copia → registro → control"),
        ("06_PLAN_NUEVO_FRONTEND_API.md", "06_Plan_Frontend_API.docx", "Plan Angular + API .NET 8", "Migración del dashboard Flask al nuevo stack"),
        ("07_MANUAL_USUARIO_WEB.md", "07_Manual_Usuario_Web.docx", "Manual de usuario — Aplicación web", "Guía con capturas de pantalla"),
        ("08_ACTA_INICIO.md", "08_Acta_Inicio_Proyecto.docx", "Acta de inicio de proyecto", "Migración de documentos a Spring"),
        ("09_ACTA_CIERRE.md", "09_Acta_Cierre_Proyecto.docx", "Acta de cierre de proyecto", "Migración de documentos a Spring"),
        ("10_CRONOGRAMA_GANTT.md", "10_Cronograma_Gantt.docx", "Cronograma y diagrama de Gantt", "Migración de adjuntos a Spring"),
        ("11_INFORME_DURACION.md", "11_Informe_Duracion.docx", "Informe de duración y fases", "Migración de documentos a Spring"),
        ("12_INFORME_CONFORMIDAD_GD.md", "12_Informe_Conformidad_GD.docx", "Informe de conformidad GD", "Jerarquía GD_Grupo y GD_Archivo"),
    ]
    for src, out_name, titulo, subt in conversiones:
        convert_md(os.path.join(DOCS, src), out_name, titulo, subt)

    convert_md(os.path.join(BASE, "FLUJO_MIGRACION_ADJUNTOS.md"),
               "13_Flujo_Migracion_Adjuntos_Detalle.docx",
               "Flujo de migración de adjuntos",
               "Portal de proveedores → Spring Gestión Documental (reglas detalladas)")

    # ----- 3. Excel de endpoints y SPs -----
    generar_excel()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

HDR_FILL = PatternFill("solid", fgColor="1F3B63")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
CELL_FONT = Font(size=10)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def generar_excel():
    wb = Workbook()

    # --- Endpoints API .NET ---
    ws = wb.active
    ws.title = "Endpoints API .NET"
    base = "api/migracion-documentos"
    rows_net = [
        ["POST", f"{base}/entornos", "—",
         "Lista los entornos configurados (Entornos_Migracion) y el entorno activo. Fallback PRUEBA si no hay BD.",
         "sp_mig_entornos", "MS_HASS_PROVEEDOR"],
        ["POST", f"{base}/cambiar-entorno", "{entorno}",
         "Cambia el entorno activo (Config_Activo), audita IP del cliente, regenera config_modo.bat / config.env. 403 si PRODUCCION no habilitada.",
         "sp_mig_entorno_activo, sp_mig_config (G)", "MS_HASS_PROVEEDOR"],
        ["POST", f"{base}/config-get", "—",
         "Devuelve la configuración del entorno activo + lista de entornos.",
         "sp_mig_entornos, sp_mig_config (G)", "MS_HASS_PROVEEDOR"],
        ["POST", f"{base}/config-update", "{updates:{k:v}}",
         "Merge sobre claves permitidas, valida (bases/raíces exactas, enums, lotes numéricos), persiste y regenera archivos.",
         "sp_mig_config (U), sp_mig_config (G)", "MS_HASS_PROVEEDOR"],
        ["POST", f"{base}/stats", "—",
         "Contadores AP/MS por estado y GATipo, GD_Archivo de prueba, última ejecución y lotes configurados.",
         "sp_mig_stats, sp_mig_gd_prueba_count", "MS_HASS_PROVEEDOR + Spring"],
        ["POST", f"{base}/archivos", "{tipo,estado,gatipo,ref,page,per_page}",
         "Bandeja paginada AP/MS. ref=OC|SO cruza claves contra AP_Documentos por lotes de 150; enriquece columna Referencias.",
         "sp_mig_archivos_ap, sp_mig_archivos_ms, sp_mig_candidatos_ap, sp_mig_claves_por_ref, sp_mig_referencias",
         "MS_HASS_PROVEEDOR + Spring"],
        ["POST", f"{base}/archivos/excel", "{tipo}",
         "Exporta a .xlsx los archivos en estado COMPLETADO, con hipervínculos a archivo y carpeta destino.",
         "sp_mig_archivos_ap, sp_mig_archivos_ms (estado COMPLETADO), sp_mig_referencias",
         "MS_HASS_PROVEEDOR + Spring"],
        ["GET", f"{base}/preview/ap/{{id}}", "—",
         "Devuelve el binario Gadatos inline (XML → text/plain).",
         "sp_mig_preview_ap", "MS_HASS_PROVEEDOR"],
        ["GET", f"{base}/preview/ms/{{id}}", "—",
         "Devuelve Contenido; si está COMPLETADO sirve el archivo desde RutaDestino.",
         "sp_mig_preview_ms", "MS_HASS_PROVEEDOR"],
        ["POST", f"{base}/log", "{tail}",
         "Últimas N líneas del log del servicio (máx. 500).",
         "(lee archivo de log)", "—"],
        ["POST", f"{base}/run", "—",
         "Ejecución manual del servicio con los lotes configurados (solo modo PRUEBA).",
         "RunManager → iniciar_servicio.bat PRUEBA", "—"],
        ["POST", f"{base}/run-all", "—",
         "Cuenta pendientes AP+MS, sube los lotes del entorno a ese total, regenera config y ejecuta.",
         "sp_mig_pendientes, sp_mig_actualizar_lotes, sp_mig_config (G)", "MS_HASS_PROVEEDOR"],
        ["POST", f"{base}/run-orden", "{numero,tipo}",
         "Valida número [0-9A-Za-z-]{1,30}; resuelve obligaciones de la OC/OS y pre-cuenta pendientes (404 si no hay); inyecta SOLO_REFERENCIA/SOLO_REFERENCIA_TIPO.",
         "sp_mig_obligaciones_referencia, sp_mig_pendientes_orden", "Spring + MS_HASS_PROVEEDOR"],
        ["POST", f"{base}/run-rango", "{periodo_desde,periodo_hasta} o {fecha_min,fecha_max}",
         "Ejecución limitada a rango de fechas o periodos YYYYMM sin tocar la config; inyecta SOLO_FECHA_MINIMA/SOLO_FECHA_MAXIMA. Permite PRUEBA y PRODUCCION habilitada.",
         "RunManager (variables de entorno al proceso)", "—"],
        ["POST", f"{base}/run-status", "{run_id}",
         "Estado y salida acumulada de una ejecución lanzada por la API.",
         "(memoria — RunManager)", "—"],
    ]
    sheet(ws, ["Método", "Endpoint", "Body / Parámetros", "Descripción",
               "SPs utilizados", "Base de datos"],
          rows_net, [9, 42, 30, 62, 44, 24])

    # --- Endpoints Flask legado ---
    ws2 = wb.create_sheet("Endpoints Flask (legado)")
    rows_flask = [
        ["GET", "/", "—", "Sirve index.html del dashboard legado.", "—"],
        ["GET", "/api/entornos", "—", "Lista Entornos_Migracion activos + entorno actual.", "Entornos_Migracion, Config_Activo"],
        ["POST", "/api/cambiar-entorno", "{entorno}", "Valida ProduccionHabilitado, actualiza Config_Activo, regenera archivos y audita.", "Config_Activo, Auditoria_Cambio_Entorno, Entornos_Migracion"],
        ["GET", "/api/config", "—", "Config del entorno activo + lista de entornos.", "Entornos_Migracion, Config_Activo"],
        ["POST", "/api/config", "{k:v}", "Valida (validate_config), persiste en Entornos_Migracion, regenera archivos.", "Entornos_Migracion"],
        ["GET", "/api/stats", "—", "Contadores AP/MS por estado y GATipo, GD_Archivo de prueba, última ejecución.", "MS_ArchivoTraslado, MS_AP_DocumentoArchivoTraslado, GD_Archivo"],
        ["GET", "/api/files", "tipo,estado,gatipo,ref,page,per_page", "Bandeja paginada; ref=OC|SO cruza claves contra AP_Documentos por lotes de 150.", "MS_Archivo, MS_AP_DocumentoArchivo, traslados, AP_Documentos"],
        ["GET", "/api/preview/ap/{id}", "—", "Binario Gadatos inline (XML → text/plain).", "MS_AP_DocumentoArchivo"],
        ["GET", "/api/preview/ms/{id}", "—", "Contenido; si COMPLETADO sirve el archivo del disco destino.", "MS_Archivo + MS_ArchivoTraslado"],
        ["GET", "/api/log?tail=N", "—", "Últimas N líneas del log.", "(archivo de log)"],
        ["POST", "/api/run", "—", "Lanza iniciar_servicio.bat PRUEBA (solo modo PRUEBA).", "—"],
        ["POST", "/api/run/all", "—", "Cuenta pendientes AP+MS, sube lotes del entorno y corre.", "traslados, Entornos_Migracion"],
        ["POST", "/api/run/orden", "{numero,tipo}", "Valida número; resuelve obligaciones y pre-cuenta pendientes (404 si no hay); lanza con SOLO_REFERENCIA.", "AP_Documentos, traslados"],
        ["GET", "/api/run/{rid}", "—", "Estado/salida de una ejecución (RunManager).", "(memoria)"],
    ]
    sheet(ws2, ["Método", "Endpoint", "Parámetros", "Descripción", "Tablas / recurso"],
          rows_flask, [9, 30, 30, 62, 46])

    # --- Stored procedures ---
    ws3 = wb.create_sheet("Stored Procedures")
    sp_headers = ["Stored procedure", "Base de datos", "Script",
                  "Descripción", "Usado por (endpoint)"]
    sps = [
        ["dbo.sp_mig_respuesta", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Helper: emite la respuesta estándar {status,message,data} FOR JSON.", "Interno (todos los SPs MS)"],
        ["dbo.sp_mig_entornos", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Lista los entornos activos de Entornos_Migracion y el entorno activo de Config_Activo.", "POST /entornos, POST /config-get"],
        ["dbo.sp_mig_entorno_activo", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "accion G/S: consulta o cambia el entorno activo; valida ProduccionHabilitado y registra Auditoria_Cambio_Entorno.", "POST /cambiar-entorno"],
        ["dbo.sp_mig_config", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "accion G/U: lee o actualiza el JSON de configuración del entorno en Entornos_Migracion.", "POST /config-get, /config-update, /cambiar-entorno, /run-all"],
        ["dbo.sp_mig_stats", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Contadores AP/MS por estado y GATipo, última ejecución registrada.", "POST /stats"],
        ["dbo.sp_mig_candidatos_ap", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Claves (id,p,t,n) de candidatos AP filtradas por estado/gatipo, para cruzar con AP_Documentos.", "POST /archivos (filtro ref=OC|SO)"],
        ["dbo.sp_mig_archivos_ap", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Página de MS_AP_DocumentoArchivo + su traslado, con filtros o por lista de ids.", "POST /archivos, /archivos/excel"],
        ["dbo.sp_mig_archivos_ms", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Página de MS_Archivo + MS_ArchivoTraslado con filtros.", "POST /archivos, /archivos/excel"],
        ["dbo.sp_mig_preview_ap", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Devuelve rowset con el binario Gadatos (no cabe en JSON).", "GET /preview/ap/{id}"],
        ["dbo.sp_mig_preview_ms", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Devuelve rowset con Contenido, Nombre, Tipo, Estado y RutaDestino.", "GET /preview/ms/{id}"],
        ["dbo.sp_mig_pendientes", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Conteo de pendientes AP + MS del entorno.", "POST /run-all"],
        ["dbo.sp_mig_pendientes_orden", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Pendientes filtrados por las obligaciones de una orden OC/OS.", "POST /run-orden"],
        ["dbo.sp_mig_actualizar_lotes", "MS_HASS_PROVEEDOR", "01_sp_ms_hass.sql",
         "Actualiza BATCH_PRUEBA / BATCH_AP_DOCUMENTO_PRUEBA del entorno.", "POST /run-all"],
        ["dbo.sp_mig_respuesta", "HASS_SPRING_PRUEBA / HP_SPRING_PRD", "02_sp_spring.sql",
         "Helper: emite la respuesta estándar {status,message,data} FOR JSON.", "Interno (SPs Spring)"],
        ["dbo.sp_mig_referencias", "HASS_SPRING_PRUEBA / HP_SPRING_PRD", "02_sp_spring.sql",
         "Devuelve ReferenciaTipo/Numero (OC/OS) de AP_Documentos para un lote de claves proveedor|tipo|num. Fallback en línea si no está desplegado.", "POST /archivos, /archivos/excel (enriquecer Referencias)"],
        ["dbo.sp_mig_claves_por_ref", "HASS_SPRING_PRUEBA / HP_SPRING_PRD", "02_sp_spring.sql",
         "Claves cuyo ReferenciaTipoDocumento coincide con el tipo pedido (OC/SO). Fallback en línea.", "POST /archivos (filtro ref)"],
        ["dbo.sp_mig_obligaciones_referencia", "HASS_SPRING_PRUEBA / HP_SPRING_PRD", "02_sp_spring.sql",
         "Obligaciones (prov,tipo,num) ligadas a un número de OC/OS. Fallback en línea.", "POST /run-orden"],
        ["dbo.sp_mig_gd_prueba_count", "HASS_SPRING_PRUEBA / HP_SPRING_PRD", "02_sp_spring.sql",
         "Cuenta GD_Archivo cuya RutaArchivo está bajo ...\\SpringGestionDoc\\PRUEBAS. Fallback en línea.", "POST /stats"],
    ]
    sheet(ws3, sp_headers, sps, [30, 30, 20, 62, 40])

    # --- Tablas ---
    ws4 = wb.create_sheet("Tablas")
    tablas = [
        ["MS_Archivo", "MS_HASS_PROVEEDOR", "Origen",
         "Adjuntos del portal de proveedores (IdArchivo, Empresa, Tabla, Id=prov|tipo|num, Nombre, Tipo, Contenido)."],
        ["MS_AP_DocumentoArchivo", "MS_HASS_PROVEEDOR", "Origen",
         "Archivos de facturas/legales (Proveedor, ObligacionTipoDocumento, ObligacionNumeroDocumento, Gatipo, Ganombre, Gadatos)."],
        ["MS_ArchivoTraslado", "MS_HASS_PROVEEDOR", "Control",
         "Estado, Intentos, Etapa, MensajeError, RutaDestino, IdGrupo, ArchivoId y fechas de cada adjunto migrado."],
        ["MS_AP_DocumentoArchivoTraslado", "MS_HASS_PROVEEDOR", "Control",
         "Mismo control para los archivos AP (facturas)."],
        ["Entornos_Migracion", "MS_HASS_PROVEEDOR", "Config dashboard",
         "Entornos (PRUEBA/PRODUCCION) con su JSON de parámetros y flag ProduccionHabilitado."],
        ["Config_Activo", "MS_HASS_PROVEEDOR", "Config dashboard",
         "Entorno activo actual."],
        ["Auditoria_Cambio_Entorno", "MS_HASS_PROVEEDOR", "Config dashboard",
         "Auditoría de cambios de entorno (usuario, IP, fecha)."],
        ["AP_Documentos", "Spring (PRUEBA/PRD)", "Origen referencial",
         "Relaciona obligación con OC/OS y da RutaArchivo (carpeta UNC nivel 6), compañía, clasificación y fecha."],
        ["GD_Grupo", "Spring (PRUEBA/PRD)", "Destino",
         "Jerarquía documental de 6 niveles: Empresa › Expediente Compras › Tipo orden › Año › Mes › Número OC/OS."],
        ["GD_Archivo", "Spring (PRUEBA/PRD)", "Destino",
         "Metadata del archivo migrado (ArchivoId, Codigo, RutaArchivo, IdGrupo, Estado)."],
        ["PersonaMast", "Spring (PRUEBA/PRD)", "Auxiliar",
         "Razón social del proveedor (DescripcionLarga del nivel 6 al crearlo)."],
        ["WH_OrdenCompra", "Spring (PRUEBA/PRD)", "Auxiliar",
         "Fecha real de la OC al crear un grupo nuevo."],
        ["Compromiso", "Spring (PRUEBA/PRD)", "Auxiliar",
         "Fecha real de la OS al crear un grupo nuevo."],
    ]
    sheet(ws4, ["Tabla", "Base de datos", "Rol", "Uso en el sistema"],
          tablas, [30, 28, 18, 80])

    out = os.path.join(OUT, "Endpoints_APIs_SPs.xlsx")
    wb.save(out)
    print("OK", "Endpoints_APIs_SPs.xlsx")


if __name__ == "__main__":
    sys.exit(main())
